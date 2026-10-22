"""Notebook export: one exam's analysis pages as Markdown / DOCX / PDF.

Page 1 is a cover (user name, exam title, date, attempt summary), then one
section per question: the question itself, the correct answer, the user's
answer, and the hand-written analysis. Deterministic, 0 token.
"""

from __future__ import annotations

import datetime as dt
import html
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from app.db import assets_dir
from app.ifmd.parser import OPTION_LETTERS
from app.repos import analyses as analyses_repo
from app.repos import documents as doc_repo
from app.repos import sessions as session_repo
from app.repos import settings as settings_repo
from app.sections import label as section_label

_OPTION_RE = re.compile(r"^([A-Z])[.)]\s*([\s\S]*)$")


@dataclass
class NotebookQuestion:
    no: int
    sec: str
    kind: str
    multi: bool
    material: str | None
    stem: str
    options: dict[str, str]
    images: list[Path] = field(default_factory=list)
    correct: list[str] = field(default_factory=list)
    chosen: list[str] = field(default_factory=list)
    is_correct: bool | None = None
    analysis: str = ""


@dataclass
class Notebook:
    title: str
    source_filename: str
    user_name: str
    generated_at: str
    session_id: int | None
    session_label: str | None
    correct: int
    graded: int
    questions: list[NotebookQuestion]

    @property
    def score(self) -> str:
        if not self.graded:
            return "—"
        return f"{self.correct}/{self.graded} ({round(self.correct / self.graded * 100)}%)"


def _option_map(options: list[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for index, raw in enumerate(options or []):
        text = (raw or "").strip()
        match = _OPTION_RE.match(text)
        if match:
            out[match.group(1)] = match.group(2).strip()
        elif index < len(OPTION_LETTERS):
            out[OPTION_LETTERS[index]] = text
    return out


def build_notebook(
    doc_id: int,
    session_id: int | None = None,
    only_analyzed: bool = False,
) -> Notebook | None:
    """Gather one exam's questions, answers, attempt and analysis pages.

    ``only_analyzed`` keeps just the questions that already have a non-empty
    analysis page (the in-app notebook preview and its exports).
    """
    doc = doc_repo.get_document(doc_id)
    if doc is None:
        return None

    questions = doc_repo.list_questions(doc_id)
    pages = {page["question_id"]: page for page in analyses_repo.list_pages(doc_id)}

    # Attempt data: the requested session, else the newest finished attempt.
    attempt: dict[int, dict] = {}
    session = None
    if session_id is not None:
        session = session_repo.get_session(session_id)
        if session is not None and session["document_id"] != doc_id:
            session = None
    if session is None:
        history = session_repo.history(doc_id, limit=20)
        finished = [h for h in history if h["finished_at"]]
        session_id = finished[0]["session_id"] if finished else None
    if session_id is not None:
        try:
            result = session_repo.regrade(session_id)
        except (KeyError, ValueError):
            result = None
        if result:
            attempt = {item["question_id"]: item for item in result["items"]}

    adir = assets_dir(doc_id)
    notebook_questions: list[NotebookQuestion] = []
    for index, q in enumerate(questions, start=1):
        item = attempt.get(q["id"], {})
        images: list[Path] = []
        for ref in q.get("images") or []:
            path = adir / str(ref).split("/")[-1]
            if path.is_file():
                images.append(path)
        notebook_questions.append(
            NotebookQuestion(
                no=q["no"] if q["no"] is not None else index,
                sec=q["sec"] or "",
                kind=q.get("kind") or "choice",
                multi=bool(q.get("multi")),
                material=q.get("material"),
                stem=q["stem"],
                options=_option_map(q.get("options") or []),
                images=images,
                correct=item.get("answers") or q.get("answers") or [],
                chosen=item.get("chosen_answers") or [],
                is_correct=item.get("is_correct"),
                analysis=(pages.get(q["id"], {}).get("content") or ""),
            )
        )

    # The cover keeps the real attempt score even when the body is filtered down
    # to the analysed pages.
    correct = sum(1 for q in notebook_questions if q.is_correct is True)
    graded = sum(1 for q in notebook_questions if q.is_correct is not None)
    if only_analyzed:
        notebook_questions = [q for q in notebook_questions if q.analysis.strip()]
    label = None
    if session_id is not None:
        label = f"Attempt #{session_id}"
        if session and session.get("started_at"):
            label += f" · {session['started_at']}"

    return Notebook(
        title=doc["title"],
        source_filename=doc.get("source_filename", ""),
        user_name=settings_repo.get_all().get("user_name", ""),
        generated_at=dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        session_id=session_id,
        session_label=label,
        correct=correct,
        graded=graded,
        questions=notebook_questions,
    )


def _answer_text(question: NotebookQuestion, values: list[str]) -> str:
    if not values:
        return "—"
    if question.kind == "blank":
        return " | ".join(values)
    return ", ".join(values)


def _option_lines(question: NotebookQuestion) -> list[str]:
    return [f"{letter}. {text}" for letter, text in question.options.items()]


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------


def render_notebook_markdown(notebook: Notebook) -> str:
    lines: list[str] = [
        f"# {notebook.title}",
        "",
        f"**Student:** {notebook.user_name or '—'}  ",
        f"**Generated:** {notebook.generated_at}  ",
    ]
    if notebook.session_label:
        lines.append(f"**Attempt:** {notebook.session_label}  ")
        lines.append(f"**Score:** {notebook.score}  ")
    lines += ["", "---", ""]

    for q in notebook.questions:
        lines.append(f"## Question {q.no} · {section_label(q.sec)}")
        lines.append("")
        if q.material:
            lines += [q.material, ""]
        lines += [q.stem, ""]
        for option in _option_lines(q):
            lines.append(f"- {option}")
        if q.options:
            lines.append("")
        lines.append(f"**Correct answer:** {_answer_text(q, q.correct)}")
        lines.append(f"**Your answer:** {_answer_text(q, q.chosen)}")
        if q.is_correct is not None:
            lines.append(f"**Result:** {'✓ correct' if q.is_correct else '✗ wrong'}")
        lines += ["", "### Analysis", "", q.analysis.strip() or "_No analysis yet._", "", "---", ""]

    return "\n".join(lines).rstrip() + "\n"


# --------------------------------------------------------------------------
# shared inline markdown (bold / italics / code / math kept verbatim)
# --------------------------------------------------------------------------


def _inline_md_to_html(text: str) -> str:
    escaped = html.escape(text)
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)
    escaped = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", escaped)
    escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
    return escaped


def _md_blocks_to_html(text: str) -> str:
    """Small markdown subset for the PDF: headings, lists, paragraphs."""
    out: list[str] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            out.append(f"<p>{'<br/>'.join(_inline_md_to_html(line) for line in paragraph)}</p>")
            paragraph.clear()

    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line.strip())
        if heading:
            flush()
            level = min(len(heading.group(1)) + 2, 6)
            out.append(f"<h{level}>{_inline_md_to_html(heading.group(2))}</h{level}>")
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            flush()
            out.append(f"<div class='bullet'>• {_inline_md_to_html(bullet.group(1))}</div>")
            continue
        paragraph.append(line.strip())
    flush()
    return "".join(out) or "<p class='muted'>No analysis yet.</p>"


def _plain_md(text: str) -> str:
    """Strip the small markdown subset for DOCX paragraphs."""
    text = re.sub(r"^#{1,4}\s+", "", text, flags=re.M)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

_CSS = """
@page {
  size: A4;
  margin: 20mm 18mm 18mm 18mm;
  @bottom-center { content: counter(page); font-size: 9pt; color: #94a3b8; }
}
* { box-sizing: border-box; }
body { font-family: "Iowan Old Style", Palatino, Georgia, serif; font-size: 10.5pt;
       line-height: 1.6; color: #1c1916; }
.cover { text-align: center; padding-top: 55mm; page-break-after: always; }
.cover .name { font-size: 22pt; letter-spacing: 0.02em; }
.cover h1 { font-size: 26pt; margin: 10mm 0 4mm; }
.cover .meta { color: #6d665a; font-size: 10pt; line-height: 2; }
.rule { border-top: 0.4mm solid #1c1916; width: 30mm; margin: 8mm auto; }
.q { page-break-before: always; }
.q-head { font-size: 9pt; letter-spacing: 0.16em; text-transform: uppercase;
          color: #6d665a; border-bottom: 0.3mm solid #dad1bf; padding-bottom: 1.5mm; }
.material { margin: 3mm 0; padding: 2mm 3mm; background: #faf8f3;
            border-left: 0.8mm solid #dad1bf; }
.stem { margin: 3mm 0; }
.opt { margin: 0.8mm 0 0.8mm 4mm; }
.answers { margin: 4mm 0; font-size: 10pt; }
.answers span { display: inline-block; min-width: 55mm; }
.good { color: #475232; font-weight: 600; }
.bad { color: #9a402c; font-weight: 600; }
.analysis { margin-top: 6mm; border-top: 0.3mm solid #dad1bf; padding-top: 3mm; }
.analysis h3 { font-size: 9pt; letter-spacing: 0.16em; text-transform: uppercase;
               color: #6d665a; margin: 0 0 2mm; }
.analysis h4, .analysis h5, .analysis h6 { font-size: 11pt; margin: 4mm 0 1mm; }
.bullet { margin-left: 4mm; }
.muted { color: #8d8678; font-style: italic; }
img { max-width: 100%; margin: 2mm 0; }
"""


def render_notebook_html(notebook: Notebook) -> str:
    """The complete, self-contained HTML source of the notebook export."""
    parts: list[str] = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<style>{_CSS}</style></head><body>",
        "<div class='cover'>",
        f"<div class='name'>{html.escape(notebook.user_name or 'Student')}</div>",
        "<div class='rule'></div>",
        f"<h1>{html.escape(notebook.title)}</h1>",
        "<div class='meta'>",
        f"Wrong-answer analysis<br/>{html.escape(notebook.generated_at)}<br/>",
    ]
    if notebook.session_label:
        parts.append(f"{html.escape(notebook.session_label)}<br/>{notebook.score}")
    parts.append("</div></div>")

    for q in notebook.questions:
        parts.append("<section class='q'>")
        parts.append(
            f"<div class='q-head'>Question {q.no} · {html.escape(section_label(q.sec))}</div>"
        )
        if q.material:
            parts.append(f"<div class='material'>{_md_blocks_to_html(q.material)}</div>")
        parts.append(f"<div class='stem'>{_md_blocks_to_html(q.stem)}</div>")
        for letter, text in q.options.items():
            parts.append(
                f"<div class='opt'><strong>{letter}.</strong> {_md_blocks_to_html(text)}</div>"
            )
        verdict = ""
        if q.is_correct is True:
            verdict = " <span class='good'>✓</span>"
        elif q.is_correct is False:
            verdict = " <span class='bad'>✗</span>"
        parts.append("<div class='answers'>")
        parts.append(
            f"<span><strong>Correct:</strong> {html.escape(_answer_text(q, q.correct))}</span>"
        )
        parts.append(
            f"<span><strong>Yours:</strong> {html.escape(_answer_text(q, q.chosen))}{verdict}</span>"
        )
        parts.append("</div>")
        parts.append(
            "<div class='analysis'><h3>Analysis</h3>"
            f"{_md_blocks_to_html(q.analysis)}</div>"
        )
        for image in q.images:
            try:
                import base64

                mime = "image/png" if image.suffix.lower() == ".png" else "image/jpeg"
                b64 = base64.b64encode(image.read_bytes()).decode("ascii")
                parts.append(f"<img src='data:{mime};base64,{b64}'/>")
            except OSError:                            # pragma: no cover
                continue
        parts.append("</section>")

    parts.append("</body></html>")
    return "".join(parts)


def render_notebook_pdf(notebook: Notebook) -> bytes:
    from app.export.pdf_helper import html_to_pdf

    return html_to_pdf(render_notebook_html(notebook))


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def render_notebook_docx(notebook: Notebook) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()

    # cover page
    for _ in range(6):
        document.add_paragraph()
    name = document.add_paragraph(notebook.user_name or "Student")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.runs[0].font.size = Pt(22)
    title = document.add_heading(notebook.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Wrong-answer analysis\n{notebook.generated_at}"
        + (f"\n{notebook.session_label}\nScore: {notebook.score}" if notebook.session_label else "")
    )
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x6D, 0x66, 0x5A)
    document.add_page_break()

    for q in notebook.questions:
        document.add_heading(f"Question {q.no} · {section_label(q.sec)}", level=2)
        if q.material:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.add_run(_plain_md(q.material)).italic = True
        document.add_paragraph(_plain_md(q.stem))
        for letter, text in q.options.items():
            document.add_paragraph(f"{letter}. {_plain_md(text)}", style="List Bullet")
        document.add_paragraph(
            f"Correct answer: {_answer_text(q, q.correct)}    "
            f"Your answer: {_answer_text(q, q.chosen)}"
            + ("    ✓" if q.is_correct is True else "    ✗" if q.is_correct is False else "")
        )
        document.add_heading("Analysis", level=3)
        document.add_paragraph(_plain_md(q.analysis) or "No analysis yet.")
        for image in q.images:
            try:
                document.add_picture(str(image), width=Inches(5.5))
            except Exception:                      # noqa: BLE001 - a bad image must not break export
                continue
        document.add_page_break()

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
