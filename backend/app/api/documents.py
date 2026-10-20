"""Document import, listing, answer-key entry and asset serving."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

from app import imports as imports_service
from app import dictionary, repos, vocab_assist
from app.db import assets_dir
from app.export import build_export, render_docx_bytes, render_pdf_bytes
from app.repos import words as words_repo
from app.ifmd.parser import OPTION_LETTERS, set_answers
from app.repos import sessions as session_repo

router = APIRouter(prefix="/api/documents", tags=["documents"])

ASSET_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")
MAX_UPLOAD = imports_service.MAX_UPLOAD
EXPORT_FORMATS = {"pdf", "docx"}

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class AnswersPayload(BaseModel):
    # ext_id -> letters (choice) or accepted strings (blank). A bare string is
    # accepted too so older clients keep working.
    answers: dict[str, list[str] | str]


class ExplainPayload(BaseModel):
    content: str = ""


class WordsPayload(BaseModel):
    """Vocabulary grid body: v2 (``columns``/``rows``/``view``) or legacy
    v1 (``headers``/``rows``). Fields are validated in the repo so a single
    ``GridError`` path yields a clean 400."""

    columns: Any = None
    rows: Any = None
    view: Any = None
    headers: Any = None


class WordEntryPayload(BaseModel):
    word: str
    meaning: str = ""
    pos: str = ""
    notes: str = ""
    context: str = ""
    source: str = ""
    question_no: int | None = None


class WordEntryPatch(BaseModel):
    word: str | None = None
    meaning: str | None = None
    pos: str | None = None
    notes: str | None = None


class SelectionPayload(BaseModel):
    text: str
    context: str = ""
    question_id: int | None = None


class EntryPayload(BaseModel):
    entry_id: int
    context: str = ""
    question_no: int | None = None


def _fail(message: str, status: int = 422) -> None:
    raise HTTPException(status_code=status, detail=message)


_SELECTION_EDGE = " \t\r\n\"'“”‘’.,;:!?()[]{}<>"


def _clean_selection(text: str) -> str:
    return " ".join(text.split()).strip(_SELECTION_EDGE)[:120]


def _question_context(doc_id: int, question_id: int | None, fallback: str):
    """Return ``(question_no, context)``; the client's context wins."""
    if question_id is None:
        return None, " ".join(fallback.split())[:2000]
    question = repos.documents.get_question(doc_id, question_id)
    if question is None:
        return None, " ".join(fallback.split())[:2000]
    parts = [question.get("material") or "", question.get("stem") or "", *question.get("options", [])]
    server_context = "\n".join(part for part in parts if part.strip())
    context = fallback.strip() or server_context
    return question.get("no"), " ".join(context.split())[:2000]


@router.post("")
async def import_document(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    filename = Path(file.filename or "upload").name
    try:
        imports_service.validate_upload(filename, data)
        result = imports_service.convert_upload(filename, data)
        doc_id = imports_service.commit(result)
    except imports_service.ImportProblem as exc:
        _fail(str(exc), exc.status)
        raise                                          # pragma: no cover

    detail = repos.documents.get_document(doc_id)
    return {**detail, "warnings": result.warnings}


@router.get("")
def list_documents() -> list[dict]:
    return repos.documents.list_documents()


@router.get("/{doc_id}")
def get_document(doc_id: int) -> dict:
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return doc


@router.get("/{doc_id}/questions")
def document_questions(doc_id: int) -> list[dict]:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return repos.documents.list_questions(doc_id)


@router.get("/{doc_id}/history")
def document_history(doc_id: int) -> list[dict]:
    return session_repo.history(doc_id)


@router.put("/{doc_id}/answers")
def put_answers(doc_id: int, payload: AnswersPayload) -> dict:
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    rows = {
        q["ext_id"]: q
        for q in repos.documents.list_questions(doc_id, include_answers=False)
    }
    cleaned: dict[str, list[str]] = {}
    for ext_id, raw in payload.answers.items():
        question = rows.get(ext_id)
        if question is None:
            continue
        values = [raw] if isinstance(raw, str) else list(raw)
        kind = question.get("kind") or "choice"
        if kind == "blank":
            accepted = [str(v).strip() for v in values if str(v).strip()]
            if not accepted:
                _fail(f"blank answer for {ext_id} must not be empty", 400)
            cleaned[ext_id] = accepted
            continue
        letters: list[str] = []
        allowed = {
            opt[0].upper()
            for opt in (question.get("options") or [])
            if opt and opt[0].isalpha()
        }
        for value in values:
            letter = str(value).strip().upper()
            if letter not in OPTION_LETTERS:
                _fail(f"invalid answer {letter!r} for {ext_id} (expected a letter A-Z)", 400)
            if allowed and letter not in allowed:
                _fail(f"answer {letter!r} has no matching option for {ext_id}", 400)
            if letter not in letters:
                letters.append(letter)
        if not letters:
            _fail(f"no answer supplied for {ext_id}", 400)
        if len(letters) > 1 and not question.get("multi"):
            _fail(
                f"{ext_id} accepts a single answer — mark it multi=1 to allow several",
                400,
            )
        cleaned[ext_id] = sorted(letters)
    if not cleaned:
        _fail("no answers supplied", 400)

    changed = repos.documents.set_answers(doc_id, cleaned)
    if changed == 0:
        _fail("no matching questions", 404)

    # keep the on-disk InteFlow-MD in sync (source of truth for export/debug)
    try:
        text = repos.documents.read_ifmd(doc_id)
        for ext_id, values in cleaned.items():
            kind = rows[ext_id].get("kind") or "choice"
            text = set_answers(text, ext_id, kind, values)
        text = re.sub(r"(?m)^answers:\s*\S+\s*$", "answers: external", text, count=1)
        repos.documents.write_ifmd(doc_id, text)
    except (KeyError, OSError):
        pass                                            # DB already updated

    return repos.documents.get_document(doc_id)


@router.put("/{doc_id}/questions/{question_id}/explain")
def put_explain(doc_id: int, question_id: int, payload: ExplainPayload) -> dict:
    """Hand-written explanation — saved straight to the DB (0 token)."""
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    content = payload.content.strip() or None
    if not repos.documents.set_explain(doc_id, question_id, content):
        _fail("question not found", 404)
        raise                                          # pragma: no cover
    return {"question_id": question_id, "explain": content}


@router.get("/{doc_id}/words")
def get_words(doc_id: int) -> dict:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return words_repo.get_grid(doc_id)


@router.put("/{doc_id}/words")
def put_words(doc_id: int, payload: WordsPayload) -> dict:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    try:
        return words_repo.save_grid(doc_id, payload.model_dump(exclude_none=True))
    except words_repo.GridError as exc:
        _fail(str(exc), 400)
        raise                                          # pragma: no cover


@router.post("/{doc_id}/words/entries")
def add_word_entry(doc_id: int, payload: WordEntryPayload) -> dict:
    """Add one word to an exam's list (de-duplicates by word)."""
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    try:
        return words_repo.add_entry(doc_id, **payload.model_dump())
    except words_repo.GridError as exc:
        _fail(str(exc), 400)
        raise                                          # pragma: no cover


@router.patch("/{doc_id}/words/entries/{row_id}")
def update_word_entry(doc_id: int, row_id: str, payload: WordEntryPatch) -> dict:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    try:
        return words_repo.update_entry(
            doc_id, row_id, payload.model_dump(exclude_unset=True)
        )
    except words_repo.EntryNotFound as exc:
        _fail(str(exc), 404)
        raise                                          # pragma: no cover
    except words_repo.GridError as exc:
        _fail(str(exc), 400)
        raise                                          # pragma: no cover


@router.delete("/{doc_id}/words/entries/{row_id}")
def delete_word_entry(doc_id: int, row_id: str) -> dict:
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    try:
        return words_repo.delete_entry(doc_id, row_id)
    except words_repo.EntryNotFound as exc:
        _fail(str(exc), 404)
        raise                                          # pragma: no cover


@router.post("/{doc_id}/words/entries/from-entry")
def add_word_from_entry(doc_id: int, payload: EntryPayload) -> dict:
    """Add the exact Merriam-Webster entry the user picked (0 token).

    Used by the Vocabulary page's dictionary lookup, where the user has already
    chosen a sense by hand — no LLM call and no sense guessing.
    """
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    try:
        entry = dictionary.get_entry(payload.entry_id)
    except dictionary.DictionaryError as exc:
        _fail(str(exc), 409)
        raise                                          # pragma: no cover
    if entry is None:
        _fail("dictionary entry not found", 404)
        raise                                          # pragma: no cover

    definition = entry["definition"] or ""
    pos = vocab_assist.first_pos(definition)
    meaning = vocab_assist.first_sense(definition)
    context = " ".join(payload.context.split())[:2000]
    try:
        result = words_repo.add_entry(
            doc_id,
            word=entry["headword"],
            meaning=meaning,
            pos=pos,
            context=context,
            source=entry["headword"],
            question_no=payload.question_no,
        )
    except words_repo.GridError as exc:
        _fail(str(exc), 400)
        raise                                          # pragma: no cover
    result.update(
        {
            "matched": True,
            "used_ai": False,
            "pos": pos,
            "definition": meaning,
            "entry_id": entry["id"],
            "dictionary_word": entry["headword"],
            "reason": "",
            "warning": "",
        }
    )
    return result


@router.post("/{doc_id}/words/entries/from-selection")
def add_word_from_selection(doc_id: int, payload: SelectionPayload) -> dict:
    """Drag-to-vocabulary: look the word up, let the LLM pick the best MW sense.

    This is an explicit user action, so an LLM call is allowed. Without a key
    (or without a matching entry) the word is still added, with a warning, so
    the gesture never dead-ends.
    """
    if repos.documents.get_document(doc_id) is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    text = _clean_selection(payload.text)
    if not text:
        _fail("select a word first", 400)
        raise                                          # pragma: no cover

    question_no, context = _question_context(
        doc_id, payload.question_id, payload.context
    )
    status = dictionary.status()
    dictionary_error = ""
    try:
        candidates = vocab_assist.candidates_for(text)
    except dictionary.DictionaryError as exc:
        candidates = []
        dictionary_error = str(exc)

    if candidates:
        choice = vocab_assist.choose_definition(text, context or text, candidates)
        chosen = next(
            (c for c in candidates if c["id"] == choice.get("entry_id")), candidates[0]
        )
        try:
            result = words_repo.add_entry(
                doc_id,
                word=chosen["headword"],
                meaning=choice["definition"],
                pos=choice["pos"],
                context=context,
                source=chosen["headword"],
                question_no=question_no,
            )
        except words_repo.GridError as exc:
            _fail(str(exc), 400)
            raise                                      # pragma: no cover
        result.update(
            {
                "matched": True,
                "used_ai": choice["used_ai"],
                "pos": choice["pos"],
                "definition": choice["definition"],
                "entry_id": chosen["id"],
                "dictionary_word": chosen["headword"],
                "reason": choice["reason"],
                "warning": choice["warning"],
                "selected": text,
            }
        )
        return result

    try:
        result = words_repo.add_entry(
            doc_id, word=text, context=context, question_no=question_no
        )
    except words_repo.GridError as exc:
        _fail(str(exc), 400)
        raise                                          # pragma: no cover
    result.update(
        {
            "matched": False,
            "used_ai": False,
            "pos": "",
            "definition": "",
            "entry_id": None,
            "dictionary_word": None,
            "reason": "",
            "warning": (
                dictionary_error
                or (
                    "Merriam-Webster is not installed; added without a definition."
                    if not status["available"]
                    else "No exact Merriam-Webster entry matched; added without a definition."
                )
            ),
            "selected": text,
        }
    )
    return result


@router.get("/{doc_id}/words/export")
def export_words(doc_id: int) -> Response:
    """Vocabulary grid as .xlsx (openpyxl; Numbers opens it too)."""
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    grid = words_repo.get_grid(doc_id)
    columns = grid["columns"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Vocabulary"
    ws.append([c["name"] for c in columns])
    for row in grid["rows"]:
        ws.append([row["cells"].get(c["id"], "") for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        # xlsx width is in characters; px→chars ≈ width / 7, clamped.
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(60, round(col["width"] / 7)))

    buf = BytesIO()
    wb.save(buf)

    filename = f"{_slugify(doc['title'])}-vocabulary.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _slugify(title: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", title).strip("-.")[:60]
    return slug or "document"


@router.get("/{doc_id}/export/{fmt}")
def export_document(doc_id: int, fmt: str) -> Response:
    """Export questions + answers + explanations as PDF or DOCX (0 token)."""
    doc = repos.documents.get_document(doc_id)
    if doc is None:
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    if fmt not in EXPORT_FORMATS:
        _fail("format must be pdf or docx", 400)
        raise                                          # pragma: no cover

    export_doc = build_export(doc_id)
    if export_doc is None:                             # pragma: no cover
        _fail("document not found", 404)
        raise
    content = render_pdf_bytes(export_doc) if fmt == "pdf" else render_docx_bytes(export_doc)
    return Response(
        content=content,
        media_type=MEDIA_TYPES[fmt],
        headers={
            "Content-Disposition": f'attachment; filename="{_slugify(doc["title"])}.{fmt}"'
        },
    )


@router.delete("/{doc_id}")
def delete_document(doc_id: int) -> dict:
    if not repos.documents.delete_document(doc_id):
        _fail("document not found", 404)
        raise                                          # pragma: no cover
    return {"deleted": doc_id}


@router.get("/{doc_id}/assets/{name}")
def document_asset(doc_id: int, name: str) -> FileResponse:
    if not ASSET_NAME_RE.match(name) or ".." in name:
        _fail("invalid asset name", 400)
        raise                                          # pragma: no cover
    path = assets_dir(doc_id) / name
    if not path.is_file():
        _fail("asset not found", 404)
        raise                                          # pragma: no cover
    media = "image/png" if name.lower().endswith(".png") else "application/octet-stream"
    return FileResponse(path, media_type=media)
